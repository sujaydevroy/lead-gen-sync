"""Tabular sources: the data.gov.in open-data API or a CSV / XLSX / JSON file downloaded from it (no AI).

Used for MCA company master data (one resource per state) and Udyam registered units. Rows are mapped with the
source's column profile, normalised, and kept only when they pass the source's filter (tobacco industry codes or
tobacco products in the activity text).
"""

from __future__ import annotations

import csv
import io
import json
from collections.abc import Iterator
from pathlib import Path
from typing import Any

from crawler.adapters.base import Adapter
from crawler.adapters.columns import column_map
from crawler.fetch import FetchError
from crawler.models import Candidate
from crawler.normalize import build_candidate

DATAGOV_API = "https://api.data.gov.in/resource/{resource}"
DATAGOV_PAGE_SIZE = 500
DATAGOV_DATASET_URL = "https://www.data.gov.in/resource/{resource}"


def read_file(path: Path) -> tuple[list[str], Iterator[list[Any]]]:
    """Headers + rows of a .csv (any common delimiter / encoding), .xlsx or .json (list of objects) file."""
    suffix = path.suffix.lower()
    if suffix in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook

        book = load_workbook(path, read_only=True, data_only=True)
        rows = book.worksheets[0].iter_rows(values_only=True)
        headers = [str(h or "").strip() for h in next(rows, [])]
        return headers, (list(r) for r in rows)
    if suffix == ".json":
        data = json.loads(path.read_text(encoding="utf-8-sig"))
        records = data.get("records", data) if isinstance(data, dict) else data
        headers = list(dict.fromkeys(k for record in records for k in record))
        return headers, ([record.get(h) for h in headers] for record in records)
    raw = path.read_bytes()
    text = raw.decode("utf-8", errors="replace")
    for encoding in ("utf-8-sig", "cp1252"):
        try:
            text = raw.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    try:
        dialect = csv.Sniffer().sniff(text[:20000], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    reader = csv.reader(io.StringIO(text), dialect)
    headers = [h.strip() for h in next(reader, [])]
    return headers, (row for row in reader)


class TabularAdapter(Adapter):
    def run(self) -> Iterator[Candidate]:
        if self.config.file:
            path = self.config.file_path()
            if not path or not path.exists():
                raise FileNotFoundError(f"{self.config.id}: file {self.config.file} not found")
            headers, rows = read_file(path)
            yield from self._candidates(headers, rows, self.config.source_url or path.name)
        else:
            yield from self._from_api()

    def _from_api(self) -> Iterator[Candidate]:
        key = self.ctx.settings.data_gov_in_api_key
        if not key:
            raise RuntimeError(
                f"{self.config.id}: set DATA_GOV_IN_API_KEY (free, data.gov.in -> My Account) or download the "
                "file and use `file:` in sources.yaml"
            )
        resource = self.config.datagov_resource
        url = DATAGOV_API.format(resource=resource)
        source_url = self.config.source_url or DATAGOV_DATASET_URL.format(resource=resource)
        offset = 0
        while True:
            params: dict[str, Any] = {"api-key": key, "format": "json", "offset": offset, "limit": DATAGOV_PAGE_SIZE}
            params.update({f"filters[{field}]": value for field, value in self.config.api_filters.items()})
            try:
                data = self.ctx.fetcher.get(url, params=params).json()
            except (FetchError, ValueError) as exc:
                self.ctx.log(f"  {self.config.id}: stopped at offset {offset}: {exc}")
                self.ctx.stats["fetch_errors"] += 1
                return
            records = data.get("records") or []
            if not records:
                return
            headers = list(dict.fromkeys(k for record in records for k in record))
            yield from self._candidates(headers, ([r.get(h) for h in headers] for r in records), source_url)
            offset += len(records)
            total = int(data.get("total") or 0)
            if (total and offset >= total) or (self.config.max_rows and offset >= self.config.max_rows):
                return

    def _candidates(self, headers: list[str], rows, source_url: str) -> Iterator[Candidate]:
        mapping = column_map(headers, self.config.profile, self.config.columns)
        if "dealer_name" not in mapping:
            raise ValueError(f"{self.config.id}: no name column among {headers[:20]}")
        for count, row in enumerate(rows, start=1):
            if self.config.max_rows and count > self.config.max_rows:
                return
            self.ctx.stats["rows_read"] += 1
            raw = {field: row[index] if index < len(row) else None for field, index in mapping.items()}
            candidate = build_candidate(
                raw,
                self.source_ref(source_url),
                extractor=f"adapter:{self.config.profile}",
                default_type=self.config.default_dealer_type,
            )
            if candidate is None:
                self.ctx.stats["rows_without_name"] += 1
                continue
            if not self.keep(self.finish(candidate)):
                continue
            self.ctx.stats["candidates"] += 1
            yield candidate
