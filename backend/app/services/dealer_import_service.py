"""Dealer file upload (system administrators): .xlsx / .xls / .csv / .json rows -> dcp.dealers.

dcp.dealers is one global DEALER DIRECTORY: a dealer belongs to no company. Clients are matched to dealers through
the products they deal in (dcp.dealer_products), never through ownership. Dealer ID is unique system-wide.

Rows are saved directly in dcp.dealers and its product tables (dcp.dealer_products, dcp.products,
dcp.product_sub_sectors); nothing else is stored about the file. created_by / modified_by on each dealer
record which system administrator added or changed it.

* Headers are matched loosely ("Dealer Name", "dealer_name", "DEALER NAME" are the same); the keys of
  dealers.json work too, and a .json file shaped like dealers.json uploads as it is. Unknown columns are
  ignored and reported.
* A row whose Dealer ID already exists UPDATES that dealer (and reactivates it); other rows
  are INSERTED. Rows without a Dealer ID always become new dealers with the next free DLR-xxxx code.
* On update only the columns present in the file change; an empty cell clears an optional value but keeps
  the required ones (name, type, status, country).
* Invalid rows are skipped and listed with their row number; valid rows are saved.
* Unknown countries, regions, sectors and dealer types are CREATED (and listed in the result); a region is
  always looked up / created within the row's country. Status and currency must be existing values.
* Every source URL is kept in dcp.dealer_sources: the row's Source URL plus an optional "sources" list
  (crawler files: JSON objects with url, kind, name, external_id, evidence, first_seen, last_seen).
"""

from __future__ import annotations

import io
import json
import re
from dataclasses import dataclass, field
from datetime import UTC, date, datetime
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from typing import Any

from sqlalchemy import func, select
from sqlalchemy.orm import Session

from app.core.errors import ApiError
from app.models import (
    Country,
    Currency,
    Dealer,
    DealerProduct,
    DealerSource,
    DealerStatus,
    DealerType,
    Product,
    ProductSubSector,
    Region,
    Sector,
    SubSector,
)
from app.schemas.admin import DealerUploadIssue, DealerUploadOut
from app.services.analytics.sector_matching import product_matches_sub_sector
from app.services.codes import highest_number
from app.services.storage_service import safe_file_name
from app.services.tabular_reader import cell_text, read_table

MAX_ISSUES_RETURNED = 1000
EMPTY_MARKERS = {"", "not available", "n/a", "na", "-", "--", "none", "null", "nil"}

# field -> (template header, accepted header spellings after normalisation)
FIELDS: dict[str, tuple[str, tuple[str, ...]]] = {
    "dealer_code": ("Dealer ID", ("dealerid", "dealercode", "dealerno", "dealernumber", "code")),
    "dealer_name": ("Dealer Name", ("dealername", "dealer", "name")),
    "legal_name": ("Company Name", ("companyname", "legalname", "registeredname", "company")),
    "dealer_type": ("Dealer Type", ("dealertype", "type")),
    "status": ("Status", ("status", "dealerstatus")),
    "contact_person": ("Contact Person", ("contactperson", "contactname", "contact")),
    "email": ("Email", ("email", "emailaddress", "emailid", "mail")),
    "phone": ("Phone", ("phone", "phonenumber", "phoneno", "mobile", "telephone", "contactnumber")),
    "website": ("Website", ("website", "websiteurl", "web", "url")),
    "registration_no": ("Registration No", ("registrationno", "registrationnumber", "regno")),
    "full_address": ("Full Address", ("fulladdress", "address")),
    "city": ("City", ("city", "town")),
    "state": ("State", ("state", "province", "stateprovince")),
    "postal_code": ("Postal Code", ("postalcode", "zipcode", "zip", "pincode", "postcode")),
    "country": ("Country", ("country",)),
    "region": ("Region", ("region",)),
    "sector": ("Sector", ("sector",)),
    "product": ("Products", ("product", "products", "productlines")),
    "last_transaction_date": ("Last Transaction Date", ("lasttransactiondate",)),
    "last_transaction_amount": ("Last Transaction Amount", ("lasttransactionamount",)),
    "currency": ("Currency", ("currency", "lasttransactioncurrency")),
    "source_url": ("Source URL", ("sourceurl", "source")),
    "verification_date": ("Verification Date", ("verificationdate", "verifiedon")),
    "is_demo": ("Is Demo", ("isdemo", "demo")),
    "created_at": ("Created At", ("createdat", "createdon", "createddate")),
    "sources": ("Sources", ("sources", "dealersources")),
}
NOT_IN_TEMPLATE = ("is_demo", "created_at", "sources")  # accepted (dealers.json has them) but not offered in the template
TEMPLATE_HEADERS = [header for name, (header, _) in FIELDS.items() if name not in NOT_IN_TEMPLATE]
TEMPLATE_EXAMPLE = [
    "", "SHAKTI POWER SYSTEMS", "Shakti Power Systems Pvt Ltd", "Distributor", "Active", "Ravi Kumar",
    "sales@shaktipower.example", "+91 98100 12345", "https://shaktipower.example", "DL-55210",
    "14 Okhla Industrial Area, New Delhi 110020, India", "New Delhi", "Delhi", "110020", "India", "North",
    "Electrical & Electrical Equipment", "MCB; Switchgear; Cables", "2026-09-15", "125000", "INR",
    "https://example.com/dealer-list", "2026-10-01",
]  # fmt: skip
MAX_LENGTHS = {
    "dealer_code": 30, "dealer_name": 200, "legal_name": 250, "contact_person": 150, "email": 254, "phone": 50,
    "website": 300, "registration_no": 100, "full_address": 500, "city": 100, "state": 100, "postal_code": 20,
    "source_url": 500,
}  # fmt: skip
TEXT_FIELDS = tuple(MAX_LENGTHS)
REQUIRED_FOR_NEW = ("dealer_name", "dealer_type", "country")
KEEP_WHEN_EMPTY = ("dealer_name", "dealer_type", "status", "country", "is_demo")
TRUE_WORDS, FALSE_WORDS = {"true", "yes", "y", "1"}, {"false", "no", "n", "0"}
COUNTRY_ALIASES = {"usa": "united states", "us": "united states", "united states of america": "united states",
                   "uk": "united kingdom", "great britain": "united kingdom", "uae": "united arab emirates"}  # fmt: skip
# Lookups an upload may create: kind -> (label in messages / results, max length)
CREATABLE = {"countries": ("Country", 100), "regions": ("Region", 50), "sectors": ("Sector", 150),
             "types": ("Dealer Type", 50)}  # fmt: skip
SOURCE_LIMITS = {"url": 500, "kind": 30, "name": 200, "external_id": 100}
DATE_FORMATS = ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%d.%m.%Y", "%Y/%m/%d", "%d %b %Y", "%d %B %Y", "%b %d, %Y")


class RowError(Exception):
    pass


def _norm(header: str) -> str:
    return re.sub(r"[^a-z0-9]", "", header.lower())


def map_columns(headers: list[str]) -> tuple[dict[str, int], list[str]]:
    """field -> column index for recognised headers (first match wins), plus the ignored headers."""
    by_alias = {alias: name for name, (_, aliases) in FIELDS.items() for alias in aliases}
    mapping: dict[str, int] = {}
    ignored: list[str] = []
    for index, header in enumerate(headers):
        name = by_alias.get(_norm(header)) if header else None
        if name and name not in mapping:
            mapping[name] = index
        elif header:
            ignored.append(header)
    return mapping, ignored


@dataclass
class _Lookups:
    countries: dict[str, int]
    regions: dict[tuple[int, str], int]  # (country id, lower-case name) -> id
    sectors: dict[str, int]
    types: dict[str, int]
    statuses: dict[str, int]
    currencies: dict[str, int]
    names: dict[str, list[str]]
    created: dict[str, list[str]] = field(default_factory=dict)  # label -> values this upload created

    @classmethod
    def load(cls, db: Session) -> _Lookups:
        def rows(model, column):
            return list(db.execute(select(column, model.id).where(model.is_active).order_by(column)))

        countries: dict[str, int] = {}
        for name, iso2, cid in db.execute(select(Country.name, Country.iso2_code, Country.id).where(Country.is_active)):
            countries[name.lower()] = cid
            if iso2:
                countries.setdefault(iso2.lower(), cid)
        for alias, target in COUNTRY_ALIASES.items():
            if target in countries:
                countries.setdefault(alias, countries[target])
        regions = {
            (country_id, name.lower()): rid
            for country_id, name, rid in db.execute(select(Region.country_id, Region.name, Region.id).where(Region.is_active))
        }
        region_names = db.scalars(
            select(Region.name).where(Region.is_active).group_by(Region.name).order_by(func.min(Region.sort_order), Region.name)
        )
        found = {
            "sectors": rows(Sector, Sector.name),
            "types": rows(DealerType, DealerType.name),
            "statuses": rows(DealerStatus, DealerStatus.name),
            "currencies": rows(Currency, Currency.code),
        }
        return cls(
            countries=countries,
            regions=regions,
            **{key: {name.lower(): rid for name, rid in values} for key, values in found.items()},
            names={"regions": list(region_names), **{key: [name for name, _ in values] for key, values in found.items()}},
        )

    def resolve(self, kind: str, label: str, value: str) -> int:
        """Id of an existing status / currency (these are never created by an upload)."""
        found = getattr(self, kind).get(value.lower())
        if found is None:
            raise RowError(f"Unknown {label} '{value}'. Use one of: {', '.join(self.names[kind])}.")
        return found

    def ensure(self, db: Session, kind: str, value: str, country_id: int | None = None) -> int:
        """Id of the country / region (within country_id) / sector / dealer type, created when it doesn't exist."""
        label, max_length = CREATABLE[kind]
        value = " ".join(value.split())
        key: Any = (country_id, value.lower()) if kind == "regions" else value.lower()
        found = getattr(self, kind).get(key)
        if found is not None:
            return found
        if len(value) > max_length:
            raise RowError(f"{label} '{value[:40]}' is longer than {max_length} characters.")
        if kind == "countries" and len(value) <= 3:
            raise RowError(f"Unknown country '{value}'. Use the full country name for a new country.")
        if kind in ("countries", "regions") and (value.islower() or value.isupper()):
            value = value.title()
        if kind == "countries":
            row: Any = Country(name=value)
        elif kind == "regions":
            last = db.scalar(select(func.max(Region.sort_order)).where(Region.country_id == country_id)) or 0
            row = Region(country_id=country_id, name=value, sort_order=last + 1)
        elif kind == "sectors":
            row = Sector(name=value)
        else:
            row = DealerType(name=value)
        db.add(row)
        db.flush()
        getattr(self, kind)[key] = row.id
        self.created.setdefault(label, []).append(value)
        return row.id


def _text(value: Any) -> str | None:
    text = cell_text(value)
    return None if text.lower() in EMPTY_MARKERS else text


def _date(value: Any, label: str) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = _text(value)
    if text is None:
        return None
    for fmt in DATE_FORMATS:
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    raise RowError(f"{label} '{text}' is not a date (use YYYY-MM-DD).")


def _amount(value: Any) -> Decimal | None:
    if isinstance(value, int | float | Decimal) and not isinstance(value, bool):
        amount = Decimal(str(value))
    else:
        text = _text(value)
        if text is None:
            return None
        try:
            amount = Decimal(re.sub(r"[,\s]", "", text))
        except InvalidOperation as exc:
            raise RowError(f"Last Transaction Amount '{text}' is not a number.") from exc
    if not amount.is_finite() or abs(amount) >= Decimal("1e16"):
        raise RowError("Last Transaction Amount is out of range.")
    return amount.quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)


def _products(value: Any) -> list[str]:
    text = _text(value)
    if text is None:
        return []
    names: list[str] = []
    for part in re.split(r"[;|,\n]+", text):
        name = part.strip()
        if name and name.lower() not in {n.lower() for n in names}:
            if len(name) > 200:
                raise RowError(f"Product '{name[:40]}...' is longer than 200 characters.")
            names.append(name)
    return names


def _flag(value: Any) -> bool | None:
    if isinstance(value, bool):
        return value
    text = _text(value)
    if text is None:
        return None
    if text.lower() in TRUE_WORDS | FALSE_WORDS:
        return text.lower() in TRUE_WORDS
    raise RowError(f"Is Demo '{text}' must be true or false.")


def parse_row(raw: dict[str, Any], lookups: _Lookups) -> dict[str, Any]:
    """Validated values for the columns present in the file (None = empty cell)."""
    values: dict[str, Any] = {}
    for name in TEXT_FIELDS:
        if name in raw:
            text = _text(raw[name])
            if text is not None and len(text) > MAX_LENGTHS[name]:
                raise RowError(f"{FIELDS[name][0]} is longer than {MAX_LENGTHS[name]} characters.")
            values[name] = text
    if values.get("email") and "@" not in values["email"]:
        raise RowError(f"Email '{values['email']}' is not a valid email address.")
    for name, (kind, label) in {"status": ("statuses", "status"), "currency": ("currencies", "currency")}.items():
        if name in raw:
            text = _text(raw[name])
            values[name] = lookups.resolve(kind, label, text) if text is not None else None
    for name in CREATABLE_FIELDS:  # kept as text here; resolve_lookups() turns them into ids (creating new ones)
        if name in raw:
            values[name] = _text(raw[name])
    for name in ("last_transaction_date", "verification_date", "created_at"):
        if name in raw:
            values[name] = _date(raw[name], FIELDS[name][0])
    if "last_transaction_amount" in raw:
        values["last_transaction_amount"] = _amount(raw["last_transaction_amount"])
    if "product" in raw:
        values["product"] = _products(raw["product"])
    if "is_demo" in raw:
        values["is_demo"] = _flag(raw["is_demo"])
    if "sources" in raw:
        values["sources"] = _sources(raw["sources"])
    return values


CREATABLE_FIELDS = {"dealer_type": "types", "country": "countries", "sector": "sectors", "region": "regions"}


def resolve_lookups(db: Session, lookups: _Lookups, values: dict[str, Any], dealer: Dealer | None) -> None:
    """Replace the text of dealer type / country / sector / region by ids, creating unknown values.

    The region is looked up within the row's country (or the dealer's current country when the row has no
    Country). When a row moves a dealer to another country without giving a region, the old region is cleared.
    """
    for name in ("dealer_type", "country", "sector"):
        if values.get(name) is not None:
            values[name] = lookups.ensure(db, CREATABLE_FIELDS[name], values[name])
    country_id = values.get("country") or (dealer.country_id if dealer is not None else None)
    if values.get("region") is not None:
        if country_id is None:
            raise RowError("Region needs a Country.")
        values["region"] = lookups.ensure(db, "regions", values["region"], country_id)
    elif "region" not in values and dealer is not None and dealer.region is not None and values.get("country"):
        if dealer.region.country_id != values["country"]:
            values["region"] = None


def _sources(value: Any) -> list[dict[str, Any]]:
    """Source list: JSON (a list of URLs or of objects with url, kind, name, external_id, evidence, first_seen,
    last_seen) or plain text with URLs separated by ; or new lines."""
    text = _text(value)
    if text is None:
        return []
    if text.startswith("["):
        try:
            items = json.loads(text)
        except ValueError as exc:
            raise RowError("Sources is not a valid JSON list.") from exc
        if not isinstance(items, list):
            raise RowError("Sources must be a list.")
    else:
        items = [part.strip() for part in re.split(r"[;\n]+", text) if part.strip()]
    sources: list[dict[str, Any]] = []
    for item in items:
        source = {"url": item} if isinstance(item, str) else item
        if not isinstance(source, dict) or not _text(source.get("url")):
            raise RowError("Every source needs a url.")
        clean: dict[str, Any] = {}
        for key, limit in SOURCE_LIMITS.items():
            text_value = _text(source.get(key))
            if text_value is not None and len(text_value) > limit:
                raise RowError(f"Source {key} '{text_value[:40]}...' is longer than {limit} characters.")
            clean[key] = text_value
        evidence = source.get("evidence")
        if evidence is not None and not isinstance(evidence, dict):
            raise RowError("Source evidence must be an object.")
        clean["evidence"] = evidence
        clean["first_seen"] = _date(source.get("first_seen"), "Source first_seen")
        clean["last_seen"] = _date(source.get("last_seen"), "Source last_seen")
        sources.append(clean)
    return sources


COLUMN_FOR = {
    "dealer_type": "dealer_type_id",
    "status": "dealer_status_id",
    "country": "country_id",
    "region": "region_id",
    "sector": "sector_id",
    "currency": "last_transaction_currency_id",
}


def _apply(dealer: Dealer, values: dict[str, Any]) -> None:
    for name, value in values.items():
        if name in ("dealer_code", "product", "created_at", "sources"):
            continue
        if value is None and name in KEEP_WHEN_EMPTY:
            continue
        setattr(dealer, COLUMN_FOR.get(name, name), value)


def _save_products(db: Session, wanted: list[tuple[Dealer, list[str]]]) -> None:
    """Replace the product lines of the given dealers (new product names are created and mapped to sub-sectors)."""
    if not wanted:
        return
    products = {p.name.lower(): p for p in db.scalars(select(Product))}
    created: list[Product] = []
    for _, names in wanted:
        for name in names:
            product = products.get(name.lower())
            if product is None:
                product = Product(name=name)
                products[name.lower()] = product
                created.append(product)
                db.add(product)
            elif not product.is_active:
                product.is_active = True
    db.flush()
    if created:
        sub_sectors = list(db.scalars(select(SubSector).where(SubSector.is_active)))
        for product in created:
            for sub_sector in sub_sectors:
                if product_matches_sub_sector(product.name, sub_sector.name):
                    db.add(ProductSubSector(product_id=product.id, sub_sector_id=sub_sector.id))

    dealer_ids = [dealer.id for dealer, _ in wanted]
    links = {
        (link.dealer_id, link.product_id): link
        for link in db.scalars(select(DealerProduct).where(DealerProduct.dealer_id.in_(dealer_ids)))
    }
    for dealer, names in wanted:
        keep = set()
        for order, name in enumerate(names, start=1):
            product_id = products[name.lower()].id
            keep.add(product_id)
            link = links.get((dealer.id, product_id))
            if link is None:
                link = DealerProduct(dealer_id=dealer.id, product_id=product_id, sort_order=order)
                links[(dealer.id, product_id)] = link
                db.add(link)
            link.is_active = True
            link.sort_order = order
        for (dealer_id, product_id), link in links.items():
            if dealer_id == dealer.id and product_id not in keep and link.is_active:
                link.is_active = False


def _save_sources(db: Session, wanted: list[tuple[Dealer, list[dict[str, Any]]]]) -> int:
    """Upsert dcp.dealer_sources by (dealer, url): first / last seen widen, given details overwrite. Returns rows saved."""
    if not wanted:
        return 0
    dealer_ids = [dealer.id for dealer, _ in wanted]
    existing = {
        (row.dealer_id, row.source_url): row
        for row in db.scalars(select(DealerSource).where(DealerSource.dealer_id.in_(dealer_ids)))
    }
    saved = 0
    for dealer, sources in wanted:
        seen_default = dealer.verification_date or date.today()
        for source in sources:
            last_seen = source["last_seen"] or seen_default
            first_seen = source["first_seen"] or last_seen
            row = existing.get((dealer.id, source["url"]))
            if row is None:
                row = DealerSource(
                    dealer_id=dealer.id, source_url=source["url"], first_seen_on=first_seen, last_seen_on=last_seen
                )
                existing[(dealer.id, source["url"])] = row
                db.add(row)
            else:
                row.first_seen_on = min(row.first_seen_on, first_seen)
                row.last_seen_on = max(row.last_seen_on, last_seen)
                row.is_active = True
            row.source_kind = source["kind"] or row.source_kind
            row.source_name = source["name"] or row.source_name
            row.external_id = source["external_id"] or row.external_id
            if source["evidence"] is not None:
                row.evidence = source["evidence"]
            saved += 1
    return saved


def _row_sources(dealer: Dealer, values: dict[str, Any]) -> list[dict[str, Any]]:
    """The row's sources list plus its Source URL (kind "upload" unless the list describes that URL)."""
    sources = list(values.get("sources") or [])
    url = values.get("source_url") if "source_url" in values else None
    if url and all(source["url"] != url for source in sources):
        blank = dict.fromkeys((*SOURCE_LIMITS, "evidence", "first_seen", "last_seen"))
        sources.append(blank | {"url": url, "kind": "upload"})
    return sources


def directory_dealer_count(db: Session) -> int:
    return db.scalar(select(func.count(Dealer.id)).where(Dealer.is_active)) or 0


def import_dealers(db: Session, *, content: bytes, file_name: str) -> DealerUploadOut:
    """Upsert the file's dealers into the dealer directory (matched by Dealer ID across all dealers)."""
    table = read_table(content, file_name)
    mapping, ignored = map_columns(table.headers)
    if "dealer_code" not in mapping and "dealer_name" not in mapping:
        raise ApiError(
            422,
            "No dealer columns were recognised in the header row. Use the template headers "
            "(Dealer ID, Dealer Name, Dealer Type, Status, Country, ...).",
        )
    lookups = _Lookups.load(db)
    existing = {d.dealer_code.upper(): d for d in db.scalars(select(Dealer))}
    next_number = max(highest_number(existing, "DLR") + 1, 1001)
    seen: set[str] = set()
    issues: list[DealerUploadIssue] = []
    product_updates: list[tuple[Dealer, list[str]]] = []
    source_updates: list[tuple[Dealer, list[dict[str, Any]]]] = []
    inserted = updated = 0

    for number, cells in table.rows:
        raw = {name: cells[index] for name, index in mapping.items()}
        try:
            values = parse_row(raw, lookups)
            code = values.get("dealer_code")
            if code and code.upper() in seen:
                raise RowError(f"Dealer ID {code} appears more than once in the file; only the first row was used.")
            dealer = existing.get(code.upper()) if code else None
            if dealer is None:
                missing = [FIELDS[name][0] for name in REQUIRED_FOR_NEW if values.get(name) is None]
                if missing:
                    raise RowError(f"New dealer is missing: {', '.join(missing)}.")
                resolve_lookups(db, lookups, values, None)
                if not code:
                    while f"DLR-{next_number}" in existing:
                        next_number += 1
                    code = f"DLR-{next_number}"
                    next_number += 1
                dealer = Dealer(dealer_code=code, is_demo=False)
                dealer.dealer_status_id = lookups.statuses.get("active")
                created = values.get("created_at")  # dealers.json "created_at"; used for new dealers only
                if created:
                    dealer.created_on = datetime(created.year, created.month, created.day, tzinfo=UTC)
                db.add(dealer)
                inserted += 1
            else:
                resolve_lookups(db, lookups, values, dealer)
                dealer.is_active = True
                updated += 1
            _apply(dealer, values)
            if dealer.dealer_status_id is None:
                raise RowError("Status is required.")
            seen.add(code.upper())
            existing[code.upper()] = dealer
            if "product" in values:
                product_updates.append((dealer, values["product"]))
            sources = _row_sources(dealer, values)
            if sources:
                source_updates.append((dealer, sources))
        except RowError as error:
            issues.append(DealerUploadIssue(row=number, message=str(error)))

    db.flush()
    _save_products(db, product_updates)
    sources_saved = _save_sources(db, source_updates)
    db.commit()
    return DealerUploadOut(
        file_name=safe_file_name(file_name, "dealers"),
        file_format=table.file_format,
        sheet_name=table.sheet_name,
        total_rows=len(table.rows),
        inserted=inserted,
        updated=updated,
        failed=len(issues),
        column_mapping={name: table.headers[index] for name, index in mapping.items()},
        issues=issues[:MAX_ISSUES_RETURNED],
        ignored_columns=ignored,
        created_lookups=lookups.created,
        sources_saved=sources_saved,
    )


def template_workbook(db: Session) -> bytes:
    """Upload template: headers + one example row, and the allowed lookup values on a second sheet."""
    from openpyxl import Workbook
    from openpyxl.styles import Font

    lookups = _Lookups.load(db)
    countries = list(db.scalars(select(Country.name).where(Country.is_active).order_by(Country.name)))
    book = Workbook()
    sheet = book.active
    sheet.title = "Dealers"
    sheet.append(TEMPLATE_HEADERS)
    sheet.append(TEMPLATE_EXAMPLE)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for column in sheet.columns:
        sheet.column_dimensions[column[0].column_letter].width = max(14, min(40, len(str(column[0].value)) + 6))

    allowed = book.create_sheet("Allowed values")
    columns = {
        "Dealer Type": lookups.names["types"],
        "Status": lookups.names["statuses"],
        "Region": lookups.names["regions"],  # each country has its own regions; these are the names in use
        "Currency": lookups.names["currencies"],
        "Country": countries,
        "Sector": lookups.names["sectors"],
    }
    allowed.append(list(columns))
    for cell in allowed[1]:
        cell.font = Font(bold=True)
    for offset in range(max(len(v) for v in columns.values())):
        allowed.append([values[offset] if offset < len(values) else None for values in columns.values()])
    for letter, width in zip("ABCDEF", (16, 12, 12, 10, 24, 48), strict=True):
        allowed.column_dimensions[letter].width = width
    notes = book.create_sheet("How to use")
    for line in (
        "Only the first sheet (Dealers) is read. Headers are matched loosely; extra columns are ignored.",
        "Dealer ID: leave empty for new dealers (a DLR-xxxx code is assigned). An existing Dealer ID updates that dealer.",
        "Required for new dealers: Dealer Name, Dealer Type, Country. Status defaults to Active.",
        "Region belongs to the row's Country. A new country, region, sector or dealer type is created automatically",
        "and listed after the upload (check the list for typos). Status and Currency must be values from 'Allowed values'.",
        "Products: separate several products with ; (semicolon).",
        "Dates: YYYY-MM-DD (or real Excel dates). Amount: a number. Currency: a code from 'Allowed values'.",
        "Also accepted: .xls (Excel 97-2003) and .csv (comma, semicolon or tab separated) with the same headers, and a",
        ".json list of dealers shaped like dealers.json (dealer_id, dealer_name, ..., product as a list, is_demo, created_at).",
    ):
        notes.append([line])
    notes.column_dimensions["A"].width = 120
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()
