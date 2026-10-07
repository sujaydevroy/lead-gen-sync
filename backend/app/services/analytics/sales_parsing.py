"""Port of src/lib/salesParsing.js + the reading part of src/services/salesService.js.

Turns the first sheet of an .xlsx workbook into normalised sales records using exactly the same
column aliases, validation rules and warning messages as the frontend.
"""

from __future__ import annotations

import io
import math
import zipfile
from dataclasses import dataclass, field
from datetime import date, datetime
from typing import Any

from openpyxl import load_workbook

SALES_COLUMNS: dict[str, list[str]] = {
    "customerName": ["customername", "customer", "dealer", "dealername"],
    "country": ["country"],
    "location": ["location", "city"],
    "month": ["month"],
    "year": ["year"],
    "date": ["date", "invoicedate", "transactiondate"],
    "product": ["product", "productname", "item"],
    "unit": ["unitofmeasurement", "uom", "unit"],
    "quantity": ["quantity", "qty"],
    "amount": ["amount", "salesamount", "value", "revenue"],
    "currency": ["currency"],
}

MONTH_NAMES = [
    "january",
    "february",
    "march",
    "april",
    "may",
    "june",
    "july",
    "august",
    "september",
    "october",
    "november",
    "december",
]


class SalesFileError(ValueError):
    """The workbook can't be used; the message is shown to the user."""


@dataclass
class ParsedSales:
    headers: list[str]
    rows: list[list[Any]]  # JSON-serialisable raw cells (header excluded)
    records: list[dict[str, Any]]
    warnings: list[dict[str, Any]] = field(default_factory=list)
    columns: dict[str, str] = field(default_factory=dict)  # recognised field -> header text
    column_indexes: dict[str, int] = field(default_factory=dict)  # recognised field -> 0-based column index


def _normalise_header(value: Any) -> str:
    return "".join(ch for ch in str(value if value is not None else "").lower() if ch.isascii() and ch.isalnum())


def serialise_cell(value: Any) -> Any:
    if isinstance(value, datetime):
        return value.isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return int(value)
    return value


def _text(value: Any) -> str:
    return "" if value is None else str(value).strip()


def _parse_month(value: Any) -> int | None:
    if value is None or value == "":
        return None
    if isinstance(value, (int, float)) and not isinstance(value, bool) and 1 <= value <= 12:
        return math.floor(value + 0.5)  # JavaScript Math.round semantics
    text = str(value).strip().lower()
    if text.isdigit() and len(text) <= 2:
        n = int(text)
        return n if 1 <= n <= 12 else None
    for index, name in enumerate(MONTH_NAMES):
        if name == text or name[:3] == text[:3]:
            return index + 1
    return None


def _parse_year(value: Any) -> int | None:
    try:
        n = value if isinstance(value, (int, float)) and not isinstance(value, bool) else float(str(value).strip())
    except (TypeError, ValueError):
        return None
    return int(n) if float(n).is_integer() and 1900 <= n <= 2200 else None


def _parse_number(value: Any) -> float | None:
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    if value is None:
        return None
    cleaned = "".join(ch for ch in str(value) if ch in "0123456789.-")
    if not cleaned or cleaned in ("-", "."):
        return None
    try:
        return float(cleaned)
    except ValueError:
        return None


def _parse_date(value: Any) -> tuple[int, int] | None:
    if not value:
        return None
    if isinstance(value, (datetime, date)):
        return value.year, value.month
    try:
        parsed = datetime.fromisoformat(str(value).replace("Z", "+00:00"))
    except ValueError:
        return None
    return parsed.year, parsed.month


def parse_sales_rows(rows: list[list[Any]]) -> ParsedSales:
    if not rows:
        raise SalesFileError("The sheet is empty.")
    header_row = rows[0]
    normalised = [_normalise_header(h) for h in header_row]
    columns: dict[str, int] = {}
    for key, aliases in SALES_COLUMNS.items():
        for index, header in enumerate(normalised):
            if header in aliases:
                columns[key] = index
                break

    missing = []
    if "customerName" not in columns:
        missing.append("CustomerName")
    if "product" not in columns:
        missing.append("Product")
    if "amount" not in columns:
        missing.append("Amount")
    if "date" not in columns and ("month" not in columns or "year" not in columns):
        missing.append("Month + Year (or Date)")
    if missing:
        found = ", ".join(str(h) for h in header_row if h) or "none"
        raise SalesFileError(f"Missing required column(s): {', '.join(missing)}. Found: {found}.")

    records: list[dict[str, Any]] = []
    warnings: list[dict[str, Any]] = []
    for i, row in enumerate(rows[1:]):
        row_number = i + 2

        def get(key: str, _row: list[Any] = row) -> Any:
            index = columns.get(key)
            return _row[index] if index is not None and index < len(_row) else None

        if all(cell is None or str(cell).strip() == "" for cell in row):
            continue

        year = _parse_year(get("year"))
        month = _parse_month(get("month"))
        if (not year or not month) and "date" in columns:
            parsed = _parse_date(get("date"))
            if parsed:
                year, month = parsed
        amount = _parse_number(get("amount"))
        customer_name = _text(get("customerName"))
        product = _text(get("product"))

        problems = []
        if not customer_name:
            problems.append("customer name is empty")
        if not product:
            problems.append("product is empty")
        if not year or not month:
            problems.append("month/year is missing or invalid")
        if amount is None:
            problems.append("amount is missing or not a number")
        if problems:
            warnings.append({"row": row_number, "message": f"Skipped: {'; '.join(problems)}."})
            continue

        currency = _text(get("currency")).upper() or "N/A"
        if currency == "N/A":
            warnings.append({"row": row_number, "message": "Currency is empty; the row is excluded from converted totals."})

        records.append(
            {
                "id": f"S-{row_number}",
                "rowNumber": row_number,
                "customerName": customer_name,
                "country": _text(get("country")) or "Unspecified",
                "location": _text(get("location")),
                "year": year,
                "month": month,
                "period": f"{year}-{month:02d}",
                "product": product,
                "unit": _text(get("unit")),
                "quantity": _parse_number(get("quantity")),
                "amount": amount,
                "currency": currency,
            }
        )

    column_names = {key: _text(header_row[index]) for key, index in columns.items()}
    return ParsedSales(
        headers=[str(h) if h is not None else "" for h in header_row],
        rows=rows[1:],
        records=records,
        warnings=warnings,
        columns=column_names,
        column_indexes=dict(columns),
    )


def read_xlsx(content: bytes, *, max_rows: int) -> list[list[Any]]:
    """Read the first worksheet as JSON-serialisable rows (trailing empty rows dropped)."""
    return read_xlsx_sheet(content, max_rows=max_rows)[1]


def read_xlsx_sheet(content: bytes, *, max_rows: int) -> tuple[str, list[list[Any]]]:
    """Like read_xlsx, but also returns the name of the worksheet that was read."""
    if not content.startswith(b"PK\x03\x04") or not zipfile.is_zipfile(io.BytesIO(content)):
        raise SalesFileError("Unsupported file type. Please upload an Excel .xlsx workbook.")
    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:  # corrupt / not a workbook
        raise SalesFileError(f"Could not read the workbook: {exc}.") from exc
    try:
        sheet = workbook.worksheets[0]
        sheet_name = sheet.title
        rows: list[list[Any]] = []
        for row in sheet.iter_rows(values_only=True):
            rows.append([serialise_cell(cell) for cell in row])
            if len(rows) > max_rows + 1:
                raise SalesFileError(f"The workbook has more than {max_rows:,} rows.")
    finally:
        workbook.close()
    while rows and all(cell is None or str(cell).strip() == "" for cell in rows[-1]):
        rows.pop()
    # Trim trailing empty columns so rows are as wide as the data.
    width = max((max((i + 1 for i, c in enumerate(r) if c not in (None, "")), default=0) for r in rows), default=0)
    return sheet_name, [r[:width] for r in rows]
