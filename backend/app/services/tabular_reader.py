"""Read the first sheet of an .xlsx / .xlsm / .xls workbook, a .csv file or a .json list into a header + rows.

The format is detected from the file content (zip = xlsx, OLE2 = xls, otherwise text = csv), so a
workbook saved with the "wrong" extension still reads. Blank rows are skipped; each row keeps its
spreadsheet row number (1 = first line of the sheet) for error messages. A .json file is a list of objects
(like dealers.json); its keys become the headers, list values are joined with "; " (lists of objects stay JSON)
and the row number is the position of the object in the list (1 = first).
"""

from __future__ import annotations

import csv
import io
import json
from dataclasses import dataclass, field
from datetime import date, datetime, time
from pathlib import PurePath
from typing import Any

from app.core.errors import ApiError

ACCEPTED_EXTENSIONS = (".xlsx", ".xlsm", ".xls", ".csv", ".json")
ZIP_MAGIC = b"PK\x03\x04"
OLE2_MAGIC = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"


@dataclass
class Table:
    file_format: str  # xlsx | xls | csv | json
    sheet_name: str | None
    headers: list[str]
    rows: list[tuple[int, list[Any]]] = field(default_factory=list)  # (row number, cells aligned with headers)


def _blank(value: Any) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def _clean(value: Any) -> Any:
    if isinstance(value, float) and value.is_integer():
        return int(value)
    if isinstance(value, str):
        return value.strip()
    return value


def _build(file_format: str, sheet_name: str | None, raw_rows, max_rows: int) -> Table:
    header: list[str] | None = None
    table = Table(file_format=file_format, sheet_name=sheet_name, headers=[])
    for number, raw in enumerate(raw_rows, start=1):
        cells = [_clean(v) for v in raw]
        if all(_blank(v) for v in cells):
            continue
        if header is None:
            header = [str(v).strip() if not _blank(v) else "" for v in cells]
            while header and not header[-1]:
                header.pop()
            table.headers = header
            continue
        if len(table.rows) >= max_rows:
            raise ApiError(422, f"The file has more than {max_rows:,} data rows. Split it into smaller files.")
        cells = (cells + [None] * len(header))[: len(header)]
        table.rows.append((number, cells))
    if not table.headers:
        raise ApiError(422, "The file is empty: no header row was found.")
    return table


def _read_xlsx(content: bytes, max_rows: int) -> Table:
    from openpyxl import load_workbook

    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:
        raise ApiError(422, "The Excel workbook could not be read. Save it again as .xlsx and retry.") from exc
    try:
        sheet = workbook.worksheets[0]
        return _build("xlsx", sheet.title, sheet.iter_rows(values_only=True), max_rows)
    finally:
        workbook.close()


def _read_xls(content: bytes, max_rows: int) -> Table:
    import xlrd

    try:
        book = xlrd.open_workbook(file_contents=content, on_demand=True)
        sheet = book.sheet_by_index(0)
    except Exception as exc:
        raise ApiError(422, "The Excel 97-2003 (.xls) workbook could not be read. Save it again and retry.") from exc

    def value(cell) -> Any:
        if cell.ctype == xlrd.XL_CELL_DATE:
            try:
                parsed = xlrd.xldate_as_datetime(cell.value, book.datemode)
            except Exception:
                return cell.value
            return parsed.date() if parsed.time() == time(0) else parsed
        if cell.ctype == xlrd.XL_CELL_BOOLEAN:
            return bool(cell.value)
        if cell.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK, xlrd.XL_CELL_ERROR):
            return None
        return cell.value

    rows = ([value(c) for c in sheet.row(i)] for i in range(sheet.nrows))
    return _build("xls", sheet.name, rows, max_rows)


def _read_csv(content: bytes, max_rows: int) -> Table:
    for encoding in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            text = content.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    if "\x00" in text:
        raise ApiError(415, "Unsupported file. Upload an Excel workbook (.xlsx, .xls) or a .csv file.")
    sample = text[:20000]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    return _build("csv", None, csv.reader(io.StringIO(text, newline=""), dialect), max_rows)


def _json_cell(value: Any) -> Any:
    if isinstance(value, list) and any(isinstance(item, dict | list) for item in value):
        return json.dumps(value, ensure_ascii=False)  # e.g. a dealer's "sources" objects
    if isinstance(value, list):
        return "; ".join(str(item) for item in value if item is not None)
    if isinstance(value, dict):
        return json.dumps(value, ensure_ascii=False)
    return _clean(value)


def _read_json(content: bytes, max_rows: int) -> Table:
    try:
        data = json.loads(content.decode("utf-8-sig"))
    except (UnicodeDecodeError, ValueError) as exc:
        raise ApiError(422, "The .json file could not be read: it is not valid JSON.") from exc
    if isinstance(data, dict):  # {"dealers": [...]} or a single dealer object
        lists = [v for v in data.values() if isinstance(v, list)]
        data = lists[0] if len(lists) == 1 else [data]
    if not isinstance(data, list) or not all(isinstance(item, dict) for item in data):
        raise ApiError(422, "The .json file must be a list of dealer objects, like dealers.json.")
    if not data:
        raise ApiError(422, "The file is empty: the .json list has no dealers.")
    if len(data) > max_rows:
        raise ApiError(422, f"The file has more than {max_rows:,} data rows. Split it into smaller files.")
    headers: list[str] = []
    for item in data:
        headers.extend(str(key) for key in item if str(key) not in headers)
    rows = [(number, [_json_cell(item.get(key)) for key in headers]) for number, item in enumerate(data, start=1)]
    return Table(file_format="json", sheet_name=None, headers=headers, rows=rows)


def read_table(content: bytes, file_name: str, *, max_rows: int = 20000) -> Table:
    extension = PurePath(file_name or "").suffix.lower()
    if extension not in ACCEPTED_EXTENSIONS:
        raise ApiError(415, "Unsupported file type. Upload an Excel workbook (.xlsx, .xls), a .csv or a .json file.")
    if not content:
        raise ApiError(422, "The file is empty.")
    if extension == ".json":
        return _read_json(content, max_rows)
    if content.startswith(ZIP_MAGIC):
        return _read_xlsx(content, max_rows)
    if content.startswith(OLE2_MAGIC):
        return _read_xls(content, max_rows)
    if extension != ".csv":
        raise ApiError(422, "The file is not a valid Excel workbook. Save it again as .xlsx and retry.")
    return _read_csv(content, max_rows)


def cell_text(value: Any) -> str:
    """A cell as display text (dates as ISO, whole numbers without .0)."""
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat() if value.time() == time(0) else value.isoformat(sep=" ")
    if isinstance(value, date):
        return value.isoformat()
    return str(_clean(value)).strip()
